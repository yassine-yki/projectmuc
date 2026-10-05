from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell


SOURCE_DIR = Path(r"C:\Users\ymerini\OneDrive - CMS\Bureau")
OUTPUT = Path(__file__).resolve().parents[1] / "src" / "equipment-data.ts"


def clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return None if not text or text in {"?", "-"} else text


def expanded_rows(path: Path, sheet_name: str):
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook[sheet_name]
    merged_values = {}
    for merged_range in sheet.merged_cells.ranges:
        value = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                merged_values[(row, column)] = value

    headers = [clean(cell.value) for cell in sheet[1]]
    result = []
    for row_number in range(2, sheet.max_row + 1):
        values = []
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row_number, column)
            value = merged_values.get((row_number, column), cell.value)
            values.append(value)
        row = {header: values[index] for index, header in enumerate(headers) if header}
        room = row.get("CHAMBER NO.")
        if isinstance(room, (int, float)) and 200 <= int(room) <= 599:
            result.append(row)
    return result


def records(path: Path, sheet: str, equipment: str, category: str, tip_column: str | None):
    by_room = {}
    for row in expanded_rows(path, sheet):
        room = int(row["CHAMBER NO."])
        product = clean(row.get("PRODUCT CODE"))
        tip = clean(row.get(tip_column)) if tip_column else product
        by_room[room] = {
            "room": room,
            "category": category,
            "equipment": equipment,
            "roomType": clean(row.get("ROOM TYPE")),
            "typology": clean(row.get("TYPOLOGIE")),
            "productCode": product,
            "tipExcel": tip,
            "source": f"{path.name} · {sheet.strip()}",
        }
    return list(by_room.values())


junior = SOURCE_DIR / "Junior.xlsx"
standard = SOURCE_DIR / "Standard.xlsx"
data = []
data += records(junior, "Headboard Junior", "headboard", "junior", "TIP EXCEL")
data += records(junior, "Armoire Junior", "wardrobe", "junior", "TIP EXCEL")
data += records(junior, "Bar Junior", "bar", "junior", None)
data += records(junior, " Vanity Junior", "vanity", "junior", None)
data += records(standard, "Headboard Standard", "headboard", "standard", "TIP EXCEL")
data += records(standard, "Armoire et Bar Standard", "wardrobe", "standard", "TIP EXCEL")
data += records(standard, "Armoire et Bar Standard", "bar", "standard", "TIP EXCEL")
data.sort(key=lambda item: (item["equipment"], item["room"]))

payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
source = (
    "// Generated from Junior.xlsx and Standard.xlsx. Run scripts/generate-equipment-data.py after workbook updates.\n"
    'export type EquipmentId = "headboard" | "wardrobe" | "bar" | "vanity";\n'
    'export type EquipmentCategory = "standard" | "junior";\n'
    "export type EquipmentRecord = { room: number; category: EquipmentCategory; equipment: EquipmentId; roomType: string | null; typology: string | null; productCode: string | null; tipExcel: string | null; source: string };\n\n"
    'export const EQUIPMENT_LABELS: Record<EquipmentId, string> = { headboard: "Headboard", wardrobe: "Armoire", bar: "Bar", vanity: "Vanity" };\n'
    f"export const EQUIPMENT_RECORDS: EquipmentRecord[] = {payload};\n\n"
    "const recordIndex = new Map(EQUIPMENT_RECORDS.map((record) => [`${record.equipment}:${record.room}`, record]));\n"
    "export function equipmentRecord(equipment: EquipmentId, room: number): EquipmentRecord | null { return recordIndex.get(`${equipment}:${room}`) || null; }\n"
    "export function equipmentIsDefined(record: EquipmentRecord | null): record is EquipmentRecord { return Boolean(record?.productCode && record?.tipExcel); }\n"
)
OUTPUT.write_text(source, encoding="utf-8")
print(f"Wrote {len(data)} records to {OUTPUT}")
